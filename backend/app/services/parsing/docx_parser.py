import io

from docx import Document


def parse_docx(content: bytes) -> tuple[str, None]:
    document = Document(io.BytesIO(content))

    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    parts.append(cell.text)

    return "\n".join(parts), None
